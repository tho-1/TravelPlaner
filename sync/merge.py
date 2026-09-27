"""Deterministic merge for sync journals (Phase 2).

Newer-wins per key, with conflicts surfaced — never silently resolved.

- Key: (store, key...). Workbook: ("workbook", destination, column).
  Trips: ("trips", trip_id, variant_id).
- last_sync.json tracks per-key applied timestamps. An entry with ts newer
  than the last applied ts for its key is a change since last sync.
- If both local and remote changed the same key since last sync to different
  values -> conflict, written to conflicts.json for UI review. Nothing is
  applied for that key until the user resolves it.
- If only one side changed -> that side wins and is applied.
- Idempotent re-sync: re-running with no new entries applies nothing.
- Offline accumulation: entries queue in local JSONL until a sync runs.
- Safety: snapshot both stores before applying anything; dry_run previews.
"""

from __future__ import annotations

import copy
import json
import os
import shutil
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def _key_of(entry: dict) -> tuple:
    return (str(entry.get("store")), *[str(k) for k in entry.get("key", [])])


def _entry_sort_key(entry: dict) -> tuple:
    return (str(entry.get("ts", "")), str(entry.get("device", "")))


def _values_equal(a, b) -> bool:
    try:
        return json.dumps(a, sort_keys=True, ensure_ascii=False) == json.dumps(
            b, sort_keys=True, ensure_ascii=False
        )
    except Exception:
        return a == b


def plan_merge(
    local_entries: list[dict],
    remote_entries: list[dict],
    last_applied: dict[str, str] | None = None,
) -> dict:
    """Pure merge planner. Returns {apply, conflicts}.

    apply: list of winning entries to apply (sorted by ts).
    conflicts: list of {key, local, remote} for keys changed on both sides
      since last sync with different values.
    """
    last_applied = last_applied or {}
    by_key: dict[tuple, dict[str, list[dict]]] = {}
    for entry in list(local_entries) + list(remote_entries):
        by_key.setdefault(_key_of(entry), {"local": [], "remote": []})
    for entry in local_entries:
        by_key[_key_of(entry)]["local"].append(entry)
    for entry in remote_entries:
        by_key[_key_of(entry)]["remote"].append(entry)

    apply: list[dict] = []
    conflicts: list[dict] = []
    for key, sides in by_key.items():
        locals_sorted = sorted(sides["local"], key=_entry_sort_key)
        remotes_sorted = sorted(sides["remote"], key=_entry_sort_key)
        key_str = json.dumps(key, ensure_ascii=False)
        since = last_applied.get(key_str)
        local_new = [e for e in locals_sorted if since is None or e["ts"] > since]
        remote_new = [e for e in remotes_sorted if since is None or e["ts"] > since]
        if local_new and remote_new:
            if _values_equal(local_new[-1]["value"], remote_new[-1]["value"]) and local_new[
                -1
            ].get("op") == remote_new[-1].get("op"):
                winner = max(
                    [local_new[-1], remote_new[-1]], key=_entry_sort_key
                )
                apply.append(winner)
            else:
                conflicts.append(
                    {"key": list(key), "local": local_new[-1], "remote": remote_new[-1]}
                )
        elif local_new:
            apply.append(local_new[-1])
        elif remote_new:
            apply.append(remote_new[-1])
    apply.sort(key=_entry_sort_key)
    return {"apply": apply, "conflicts": conflicts}


def _snapshot_file(path: Path, backup_dir: Path, prefix: str) -> Path | None:
    try:
        if not path.exists():
            return None
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%Y%m%d-%H%M%S")
        millis = int(time.time() * 1000) % 1000
        target = backup_dir / f"{prefix}-{stamp}-{millis:03d}{path.suffix}"
        shutil.copy2(path, target)
        return target
    except OSError:
        return None


def apply_entries(
    entries: list[dict],
    workbook_path: Path | None = None,
    trips_path: Path | None = None,
    dry_run: bool = False,
) -> dict:
    """Apply winning entries to both stores. Snapshots first. Returns summary.

    Sets SYNC_MERGE_APPLY=1 so the apply itself is not re-journaled.
    """
    from environment import get_workbook_path

    wb_path = Path(workbook_path) if workbook_path else get_workbook_path()
    tr_path = Path(trips_path) if trips_path else (ROOT / "trips.json")
    snapshots = []
    if dry_run:
        return {"would_apply": len(entries), "snapshots": [], "applied": 0}
    for path, backup_dir, prefix in (
        (wb_path, wb_path.parent / "workbook_backups", wb_path.stem),
        (tr_path, tr_path.parent / "trips_backups", "trips"),
    ):
        snap = _snapshot_file(path, backup_dir, prefix)
        if snap is not None:
            snapshots.append(str(snap))
    os.environ["SYNC_MERGE_APPLY"] = "1"
    try:
        applied = 0
        for entry in sorted(entries, key=_entry_sort_key):
            if _apply_one(entry, wb_path, tr_path):
                applied += 1
    finally:
        os.environ.pop("SYNC_MERGE_APPLY", None)
    return {"applied": applied, "snapshots": snapshots}


def _apply_one(entry: dict, wb_path: Path, tr_path: Path) -> bool:
    try:
        if entry.get("store") == "workbook":
            return _apply_workbook(entry, wb_path)
        if entry.get("store") == "trips":
            return _apply_trips(entry, tr_path)
    except Exception:
        return False
    return False


def _apply_workbook(entry: dict, wb_path: Path) -> bool:
    import openpyxl

    from data_utils import (
        _find_destination_sheet,
        load_workbook_for_update,
        save_workbook_atomic,
    )

    key = entry.get("key", [])
    if len(key) != 2:
        return False
    destination, column = str(key[0]), str(key[1])
    value = entry.get("value")
    sheet = _find_destination_sheet(wb_path)
    if sheet is None:
        return False
    wb = load_workbook_for_update(wb_path)
    try:
        ws = wb[sheet]
        headers = {}
        for idx in range(1, ws.max_column + 1):
            header_value = ws.cell(row=1, column=idx).value
            if header_value is not None:
                headers[str(header_value)] = idx
        dest_idx = next(
            (i for h, i in headers.items() if "destination" in h.lower()), None
        )
        col_idx = headers.get(column)
        if col_idx is None:
            lowered = column.strip().lower()
            col_idx = next(
                (i for h, i in headers.items() if h.strip().lower() == lowered),
                None,
            )
        if dest_idx is None or col_idx is None:
            wb.close()
            return False
        row_idx = next(
            (
                r
                for r in range(2, ws.max_row + 1)
                if ws.cell(r, dest_idx).value is not None
                and str(ws.cell(r, dest_idx).value).strip().lower()
                == destination.strip().lower()
            ),
            None,
        )
        if row_idx is None:
            wb.close()
            return False
        ws.cell(row=row_idx, column=col_idx).value = copy.deepcopy(value)
        save_workbook_atomic(wb, wb_path)
        wb.close()
        return True
    except Exception:
        try:
            wb.close()
        except Exception:
            pass
        return False


def _apply_trips(entry: dict, tr_path: Path) -> bool:
    from itinerary import storage

    key = entry.get("key", [])
    if len(key) != 2:
        return False
    trip_id, variant_id = str(key[0]), str(key[1])
    data = storage.load_from(tr_path)
    trips = {t["id"]: t for t in data.get("trips", [])}
    changed = False
    if entry.get("op") == "delete":
        trip = trips.get(trip_id)
        if trip is not None:
            before = len(trip.get("variants", []))
            trip["variants"] = [
                v for v in trip.get("variants", []) if v.get("id") != variant_id
            ]
            if len(trip["variants"]) != before:
                changed = True
            if not trip["variants"]:
                data["trips"] = [t for t in data.get("trips", []) if t["id"] != trip_id]
                changed = True
    else:
        variant = copy.deepcopy(entry.get("value"))
        if not isinstance(variant, dict):
            return False
        variant["id"] = variant_id
        trip = trips.get(trip_id)
        if trip is None:
            return False
        for i, existing in enumerate(trip.get("variants", [])):
            if existing.get("id") == variant_id:
                if not _values_equal(existing, variant):
                    trip["variants"][i] = variant
                    changed = True
                break
        else:
            trip["variants"].append(variant)
            changed = True
    if changed:
        storage.save_to(data, tr_path)
    return changed


def load_last_applied(journal_dir: Path | str | None = None) -> dict[str, str]:
    from sync import journal as _journal

    jdir = Path(journal_dir) if journal_dir else _journal.DEFAULT_JOURNAL_DIR
    path = jdir / "last_sync.json"
    try:
        return json.loads(path.read_text(encoding="utf-8")).get("applied", {})
    except Exception:
        return {}


def save_last_applied(
    applied_entries: list[dict],
    journal_dir: Path | str | None = None,
) -> None:
    from sync import journal as _journal

    jdir = Path(journal_dir) if journal_dir else _journal.DEFAULT_JOURNAL_DIR
    jdir.mkdir(parents=True, exist_ok=True)
    current = load_last_applied(jdir)
    for entry in applied_entries:
        current[json.dumps(_key_of(entry), ensure_ascii=False)] = str(entry["ts"])
    (jdir / "last_sync.json").write_text(
        json.dumps(
            {"last_sync_utc": _journal._utcnow_iso(), "applied": current},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def load_conflicts(journal_dir: Path | str | None = None) -> list[dict]:
    from sync import journal as _journal

    jdir = Path(journal_dir) if journal_dir else _journal.DEFAULT_JOURNAL_DIR
    try:
        return json.loads((jdir / "conflicts.json").read_text(encoding="utf-8"))
    except Exception:
        return []


def save_conflicts(conflicts: list[dict], journal_dir: Path | str | None = None) -> None:
    from sync import journal as _journal

    jdir = Path(journal_dir) if journal_dir else _journal.DEFAULT_JOURNAL_DIR
    jdir.mkdir(parents=True, exist_ok=True)
    (jdir / "conflicts.json").write_text(
        json.dumps(conflicts, ensure_ascii=False, indent=2), encoding="utf-8"
    )
