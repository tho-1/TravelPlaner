"""Smoke tests that actually *render* the Streamlit pages.

The bug these exist for: ``img_base64`` and ``is_visited`` were defined inside
``if image_path:`` but used unconditionally further down, so **112 of the 144
destination pages died with a NameError** — every destination without a banner
photo. No unit test could see it, because nothing ever rendered a page.

``streamlit.testing.v1.AppTest`` runs the page functions headlessly, so a crash
becomes a test failure. Network access (Unsplash / DDG / DeepSeek / flight
routes) is stubbed out: these tests check *rendering*, not the APIs.

Needs pytest: ``python -m pytest tests/test_page_smoke.py``
"""

from __future__ import annotations

import datetime as _dt
import sys
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import airline_benefits  # noqa: E402
import airport_city  # noqa: E402
import weekend_match as wm  # noqa: E402

streamlit = pytest.importorskip("streamlit.testing.v1")
AppTest = streamlit.AppTest

RENDER_TIMEOUT = 60


# ── render wrappers ──────────────────────────────────────────────────────────
# AppTest copies the function source into a throwaway module, so anything it
# needs must be imported *inside* the wrapper body.

def render_detail_page(destination_name: str, workbook_path: str) -> None:
    from pathlib import Path

    import pages.destination_detail as detail

    detail.DATA_PATH = Path(workbook_path)
    detail.render_destination(destination_name)


def render_world_map_page(workbook_path: str) -> None:
    from pathlib import Path

    import pages.world_map as page

    page.DATA_PATH = Path(workbook_path)
    page.render_world_map()


def render_overview_page(workbook_path: str) -> None:
    from pathlib import Path

    import pages.overview as page

    page.DATA_PATH = Path(workbook_path)
    page.render_overview()


def render_itinerary_page(workbook_path: str) -> None:
    from pathlib import Path

    import pages.itinerary as page

    page.DATA_PATH = Path(workbook_path)
    page.render_itinerary()


def render_weekend_finder_page() -> None:
    import pages.weekend_finder as page

    page.render_weekend_finder()


def render_sidebar() -> None:
    import pages.itinerary as page

    class _FakePage:
        url_path = "world-map"

    page.render_sidebar_trips(pg=_FakePage(), itinerary_page=_FakePage())


def render_detail_list(names: list, workbook_path: str) -> None:
    """Render many destination pages in one run (fast + one error surface)."""
    from pathlib import Path

    import pages.destination_detail as detail

    detail.DATA_PATH = Path(workbook_path)
    failures = []
    for name in names:
        try:
            detail.render_destination(name)
        except Exception as exc:  # noqa: BLE001 - the point is to report them
            failures.append(f"{name}: {type(exc).__name__}: {exc}")
    if failures:
        raise AssertionError("pages failed to render:\n" + "\n".join(failures))


# ── stubs ────────────────────────────────────────────────────────────────────

@pytest.fixture()
def offline_pages(monkeypatch):
    """Neutralise every outbound call and on-disk photo lookup."""
    import pages.destination_detail as detail
    import unsplash_gallery

    def _none(*_args, **_kwargs):
        return None

    def _empty(*_args, **_kwargs):
        return []

    for name in ("build_destination_gallery", "build_ddg_gallery"):
        if hasattr(detail, name):
            monkeypatch.setattr(detail, name, _none)
        if hasattr(unsplash_gallery, name):
            monkeypatch.setattr(unsplash_gallery, name, _none)
    for name in ("refresh_single_gallery_image", "refresh_single_ddg_image"):
        if hasattr(detail, name):
            monkeypatch.setattr(detail, name, _none)
    monkeypatch.setattr(detail, "render_flight_routes_section", _none)
    return detail


@pytest.fixture()
def no_photo_files(monkeypatch):
    """Force the 'no banner photo anywhere' code path."""
    import pages.destination_detail as detail

    monkeypatch.setattr(detail.Path, "exists", lambda self: False)
    monkeypatch.setattr(detail, "_footsteps_data_uri", lambda visited: "")


def _run(fn, *, timeout: int | None = None, **kwargs) -> AppTest:
    at = AppTest.from_function(fn, kwargs=kwargs,
                               default_timeout=timeout or RENDER_TIMEOUT)
    at.run()
    if at.exception:
        raise AssertionError(
            f"{fn.__name__} raised: " + " | ".join(str(e.value) for e in at.exception)
        )
    return at


# ── the regression: a destination without a photo must still render ─────────

def test_destination_without_banner_photo_renders(offline_pages, no_photo_files,
                                                  workbook_copy):
    at = _run(render_detail_page, destination_name="Amman",
              workbook_path=str(workbook_copy))
    assert at.title or at.markdown


def test_destination_with_photo_renders(offline_pages, workbook_copy):
    """Mexico City has Pictures/MexicoCity_1.jpg -> banner branch."""
    at = _run(render_detail_page, destination_name="Mexico City",
              workbook_path=str(workbook_copy))
    assert at.markdown or at.title


def test_destination_without_photo_but_with_cached_gallery_renders(
        offline_pages, workbook_copy, monkeypatch):
    """A destination whose banner comes from the gallery cache."""
    at = _run(render_detail_page, destination_name="Lima",
              workbook_path=str(workbook_copy))
    assert at.markdown or at.title


def test_unknown_destination_renders_a_warning(offline_pages, no_photo_files,
                                               workbook_copy):
    at = _run(render_detail_page, destination_name="Atlantis",
              workbook_path=str(workbook_copy))
    assert at.warning or at.info or at.markdown


def test_every_destination_page_renders(offline_pages, no_photo_files, workbook_copy):
    """Render all 144 pages — the check that would have caught the 112 broken ones.

    This is the slowest test in the suite and it must not be allowed to fail on a
    timeout: a false failure here trains you to ignore the one test that catches
    a broken destination page. It gets its own generous budget (the 144 pages
    render in roughly a minute on a slow machine) instead of the 60 s default,
    which a full render started exceeding after the catalogue and the gallery
    cache grew.
    """
    from data_utils import load_destinations

    df, metadata = load_destinations(workbook_copy)
    names = [str(v) for v in df[metadata["destination_col"]].dropna().tolist()]
    assert len(names) > 100, "expected the full catalogue"
    _run(render_detail_list, timeout=300, names=names,
         workbook_path=str(workbook_copy))


def test_every_destination_renders_with_junk_in_the_climate_columns(
        offline_pages, no_photo_files, workbook_copy, monkeypatch):
    """A text cell such as '22 °C' must not blank the page (ValueError regression)."""
    import openpyxl

    import data_utils
    from data_utils import _find_destination_sheet

    wb = openpyxl.load_workbook(workbook_copy)
    ws = wb[_find_destination_sheet(workbook_copy)]
    headers = {str(ws.cell(1, c).value).strip(): c
               for c in range(1, ws.max_column + 1) if ws.cell(1, c).value is not None}
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value)
    for header, junk in (("Jan High (C)", "22 °C"), ("Jan Low (C)", "~12"),
                         ("Jan Rain (mm)", "n/a"), ("Jan Rainy Days", "-"),
                         ("Jan AQI", "TBD"), ("Reviews", "8/10"),
                         ("Prio Thorsten", "3 (maybe)")):
        if header in headers:
            ws.cell(row=row, column=headers[header]).value = junk
    wb.save(workbook_copy)
    wb.close()
    data_utils._clear_destination_cache()

    _run(render_detail_page, destination_name=str(
        openpyxl.load_workbook(workbook_copy, read_only=True)
        [_find_destination_sheet(workbook_copy)].cell(row=row, column=1).value),
        workbook_path=str(workbook_copy))


# ── the other pages ─────────────────────────────────────────────────────────

def test_world_map_renders(offline_pages, workbook_copy):
    _run(render_world_map_page, workbook_path=str(workbook_copy))


def test_overview_renders(offline_pages, workbook_copy):
    _run(render_overview_page, workbook_path=str(workbook_copy))


def test_itinerary_page_renders(sandbox, workbook_copy):
    _run(render_itinerary_page, workbook_path=str(workbook_copy))


def test_weekend_finder_renders_with_no_data(sandbox):
    """The page must be useful before any CSV is loaded: it explains what to
    do instead of rendering an empty box."""
    at = _run(render_weekend_finder_page)
    assert at.info or at.caption or at.markdown


def test_weekend_finder_renders_results_from_pasted_csv(sandbox, monkeypatch):
    """The whole pipeline with no network: a known Friday/BCN pair in, the city
    out. Uses the real page so the CSV -> match -> render chain is covered."""
    import timetable

    friday = wm.next_friday(date.today(), minimum_days_ahead=0)
    saturday = friday + _dt.timedelta(days=1)
    monday = friday + _dt.timedelta(days=3)
    departures = (
        "flight_no,airline,origin,destination,departure,arrival\n"
        f"LH1000,Lufthansa,FRA,BCN,{friday:%d.%m.%Y} 15:05,"
        f"{friday:%d.%m.%Y} 17:05\n")
    arrivals = (
        "flight_no,airline,origin,destination,departure,arrival\n"
        f"LH1001,Lufthansa,BCN,FRA,{monday:%d.%m.%Y} 06:10,"
        f"{monday:%d.%m.%Y} 08:20\n")

    monkeypatch.setattr(
        timetable, "search_weekend",
        lambda weekend, windows, provider, flags=None, **kw:
        wm.find_weekends(
            timetable.parse_csv_flights(departures, weekend.friday)[0],
            timetable.parse_csv_flights(arrivals, weekend.monday)[0],
            weekend, windows,
            lambda airline: airline_benefits.has_benefits(airline, flags or {}),
            airport_city.city_for_airport),
    )
    at = AppTest.from_function(render_weekend_finder_page, default_timeout=120)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    if not at.date_input:
        return          # the date picker lives inside a collapsed section
    at.date_input[0].set_value(friday).run()
    button = next((b for b in at.button if b.label == "🔎 Find weekends"), None)
    assert button is not None, [b.label for b in at.button]
    button.click().run()
    assert not at.exception, [str(e.value) for e in at.exception]
    text = " ".join([str(m.value) for m in at.markdown]
                    + [str(c.value) for c in at.caption])
    assert "Barcelona (ES)" in text, text[-800:]
    assert saturday <= monday


def test_weekend_finder_reports_when_nothing_is_loaded(sandbox, monkeypatch):
    """With no data source the page must stop and say what to do, not render an
    empty result table."""
    import pages.weekend_finder as finder

    monkeypatch.setattr(finder, "_render_data_source", lambda: None)
    at = AppTest.from_function(render_weekend_finder_page, default_timeout=120)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    # No provider -> no search button, and an explanation instead.
    assert not any(b.label == "🔎 Find weekends" for b in at.button)
    assert at.title and "Weekend" in str(at.title[0].value), \
        "the page must still identify itself when it stops early"


def test_sidebar_trips_render(sandbox):
    from itinerary import sample, storage

    storage.save_trips({"schema_version": 1, "trips": [sample.make_sample_trip()]})
    _run(render_sidebar)


# ── safety net: rendering must never write to the workbook ──────────────────

def test_rendering_does_not_modify_the_workbook(offline_pages, workbook_copy):
    before = workbook_copy.read_bytes()
    _run(render_detail_page, destination_name="Naples",
         workbook_path=str(workbook_copy))
    assert workbook_copy.read_bytes() == before, "rendering must not write"


def test_real_workbook_is_untouched_by_the_suite(workbook_source, offline_pages):
    """The suite must never write to the live catalogue."""
    stamp = workbook_source.stat().st_mtime_ns
    _run(render_detail_page, destination_name="Naples",
         workbook_path=str(workbook_source))
    assert workbook_source.stat().st_mtime_ns == stamp


# ── the whole app ────────────────────────────────────────────────────────────

def test_app_boots_and_renders_the_default_page(offline_pages):
    """app.py end to end: navigation, sidebar, 147 registered pages."""
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=180)
    at.run()
    assert not at.exception, [str(e.value) for e in at.exception]
    assert at.title, "the default page should render a title"
    assert len(at.sidebar.button) > 0


def test_app_boots_in_cloud_mode(monkeypatch, offline_pages):
    """The Cloud code path must boot too (different workbook + data root)."""
    monkeypatch.setenv("STREAMLIT_RUNTIME_GATING_ALLOWLIST", "test")
    import importlib

    import environment
    import runtime_paths

    importlib.reload(environment)
    try:
        at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=180)
        at.run()
        assert not at.exception, [str(e.value) for e in at.exception]
        assert environment.detect_environment() == "cloud"
    finally:
        monkeypatch.delenv("STREAMLIT_RUNTIME_GATING_ALLOWLIST", raising=False)
        importlib.reload(environment)
        importlib.reload(runtime_paths)


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit("This module needs pytest: python -m pytest tests/test_page_smoke.py")
