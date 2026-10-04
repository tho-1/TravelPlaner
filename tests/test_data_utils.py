from __future__ import annotations

import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import data_utils


def _make_workbook(path: Path, value: str = "base") -> None:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = value
    workbook.save(path)
    workbook.close()


def _read_value(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, read_only=True)
    value = workbook.active["A1"].value
    workbook.close()
    return value


def test_atomic_save_snapshots_outgoing_workbook():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.xlsx"
        _make_workbook(path)
        workbook = data_utils.load_workbook_for_update(path)
        workbook.active["A1"] = "updated"
        data_utils.save_workbook_atomic(workbook, path)
        workbook.close()

        backups = list((Path(td) / "workbook_backups").glob("sample-*.xlsx"))
        assert len(backups) == 1
        assert _read_value(backups[0]) == "base"
        assert _read_value(path) == "updated"


def test_update_comment_uses_atomic_writer():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Destination", "Comment"])
        sheet.append(["Example", None])
        workbook.save(path)
        workbook.close()

        data_utils.update_comment("Example", "saved through writer", path=path)

        check = openpyxl.load_workbook(path, read_only=True)
        assert check.active["B2"].value == "saved through writer"
        check.close()
        backups = list((Path(td) / "workbook_backups").glob("sample-*.xlsx"))
        assert len(backups) == 1
        assert _read_value_from_cell(backups[0], "B2") is None


def _read_value_from_cell(path: Path, cell: str):
    workbook = openpyxl.load_workbook(path, read_only=True)
    value = workbook.active[cell].value
    workbook.close()
    return value


def test_stale_workbook_cannot_overwrite_newer_save():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.xlsx"
        _make_workbook(path)
        stale = data_utils.load_workbook_for_update(path)
        current = data_utils.load_workbook_for_update(path)
        stale.active["A1"] = "stale"
        current.active["A1"] = "current"
        data_utils.save_workbook_atomic(current, path)
        current.close()

        try:
            data_utils.save_workbook_atomic(stale, path)
        except data_utils.WorkbookChangedError:
            pass
        else:
            raise AssertionError("stale workbook save was not rejected")
        finally:
            stale.close()
        assert _read_value(path) == "current"


def test_workbook_backups_keep_newest_twelve():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.xlsx"
        _make_workbook(path)
        for value in range(14):
            workbook = data_utils.load_workbook_for_update(path)
            workbook.active["A1"] = str(value)
            data_utils.save_workbook_atomic(workbook, path)
            workbook.close()

        backups = sorted((Path(td) / "workbook_backups").glob("sample-*.xlsx"))
        assert len(backups) == data_utils.WORKBOOK_BACKUP_KEEP
        assert _read_value(backups[0]) == "1"
        assert _read_value(backups[-1]) == "12"
        assert _read_value(path) == "13"


def test_transient_lock_is_retried_and_a_permanent_one_is_not():
    """Antivirus/indexer holds a file for a moment; Excel holds it for good.

    The difference must be *retries*, not luck, because both surface to
    openpyxl as the same PermissionError.
    """
    attempts = []

    def flaky():
        attempts.append(1)
        if len(attempts) < 3:
            raise PermissionError("scanner")
        return "saved"

    assert data_utils.retry_while_locked(flaky, retries=5, delay=0) == "saved"
    assert len(attempts) == 3

    attempts.clear()

    def locked_forever():
        attempts.append(1)
        raise PermissionError("open in Excel")

    try:
        data_utils.retry_while_locked(locked_forever, retries=4, delay=0)
    except PermissionError:
        pass
    else:
        raise AssertionError("a permanent lock must still surface")
    assert len(attempts) == 4, "it must stop after the configured retries"


def test_save_survives_a_transient_replace_denial():
    """The flake this retries: Windows denies the replace once, then allows it."""
    import os as _os

    real_replace = _os.replace
    calls = []

    def deny_once(src, dst):
        calls.append(1)
        if len(calls) == 1:
            raise PermissionError("transient")
        return real_replace(src, dst)

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.xlsx"
        _make_workbook(path)
        workbook = data_utils.load_workbook_for_update(path)
        workbook.active["A1"] = "eventually"
        try:
            with patch("data_utils.os.replace", side_effect=deny_once), \
                    patch("data_utils.time.sleep"):
                data_utils.save_workbook_atomic(workbook, path)
        finally:
            workbook.close()
        assert _read_value(path) == "eventually"
        assert len(calls) == 2


def test_replace_failure_preserves_original_and_raises_friendly_error():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "sample.xlsx"
        _make_workbook(path)
        workbook = data_utils.load_workbook_for_update(path)
        workbook.active["A1"] = "not installed"
        try:
            # sleep is patched out: the retry backoff must not slow the tests
            with patch("data_utils.os.replace", side_effect=PermissionError("locked")), \
                    patch("data_utils.time.sleep"):
                try:
                    data_utils.save_workbook_atomic(workbook, path)
                except data_utils.WorkbookLockedError:
                    pass
                else:
                    raise AssertionError("replace failure did not raise a friendly error")
        finally:
            workbook.close()
        assert _read_value(path) == "base"
        assert list((Path(td) / "workbook_backups").glob("sample-*.xlsx"))


if __name__ == "__main__":
    tests = [
        test_atomic_save_snapshots_outgoing_workbook,
        test_update_comment_uses_atomic_writer,
        test_stale_workbook_cannot_overwrite_newer_save,
        test_workbook_backups_keep_newest_twelve,
        test_transient_lock_is_retried_and_a_permanent_one_is_not,
        test_save_survives_a_transient_replace_denial,
        test_replace_failure_preserves_original_and_raises_friendly_error,
    ]
    for test in tests:
        test()
        print(f"PASS {test.__name__}")
    print(f"{len(tests)}/{len(tests)} passed")