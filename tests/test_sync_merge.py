from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from itinerary import storage
from sync import journal, merge


def _make_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Destinations"
    ws.cell(1, 1, "Destination")
    ws.cell(1, 2, "Comment")
    ws.cell(2, 1, "Naples")
    ws.cell(2, 2, "base")
    wb.save(path)
    wb.close()


def _make_trips(path: Path) -> dict:
    data = {
        "schema_version": 1,
        "trips": [
            {
                "id": "trip-1",
                "name": "T",
                "created": "2026-01-01T00:00:00",
                "active_variant_id": "variant-1",
                "variants": [
                    {
                        "id": "variant-1",
                        "name": "Variant 1",
                        "notes": "",
                        "rating": None,
                        "months": [],
                        "comment": "",
                        "stops": [],
                        "legs": [],
                    }
                ],
            }
        ],
    }
    storage.save_to(data, path)
    return storage.load_from(path)


def _entry(store, key, value, ts, device, op="upsert"):
    return {
        "ts": ts,
        "device": device,
        "store": store,
        "op": op,
        "key": key,
        "value": value,
    }


def test_newer_wins_without_conflict():
    remote = [_entry("workbook", ["Naples", "Comment"], "new", "2026-09-27T11:00:00+00:00", "cloud")]
    plan = merge.plan_merge([], remote, {})
    assert len(plan["conflicts"]) == 0
    assert len(plan["apply"]) == 1
    assert plan["apply"][0]["value"] == "new"
    # Both sides agree on the value -> no conflict, single apply.
    local = [_entry("workbook", ["Naples", "Comment"], "same", "2026-09-27T10:00:00+00:00", "local")]
    remote2 = [_entry("workbook", ["Naples", "Comment"], "same", "2026-09-27T11:00:00+00:00", "cloud")]
    plan2 = merge.plan_merge(local, remote2, {})
    assert len(plan2["conflicts"]) == 0
    assert len(plan2["apply"]) == 1


def test_conflict_when_both_sides_changed():
    last = {json.dumps(("workbook", "Naples", "Comment")): "2026-09-27T09:00:00+00:00"}
    local = [_entry("workbook", ["Naples", "Comment"], "local-edit", "2026-09-27T10:00:00+00:00", "local")]
    remote = [_entry("workbook", ["Naples", "Comment"], "cloud-edit", "2026-09-27T10:30:00+00:00", "cloud")]
    plan = merge.plan_merge(local, remote, last)
    assert len(plan["conflicts"]) == 1
    assert len(plan["apply"]) == 0
    assert plan["conflicts"][0]["key"] == ["workbook", "Naples", "Comment"]


def test_resolution_and_idempotent_resync():
    with tempfile.TemporaryDirectory() as td:
        jdir = Path(td) / "journals"
        jdir.mkdir()
        wb = Path(td) / "wb.xlsx"
        trips = Path(td) / "trips.json"
        _make_workbook(wb)
        _make_trips(trips)
        winner = _entry(
            "workbook", ["Naples", "Comment"], "resolved",
            "2026-09-27T12:00:00+00:00", "local",
        )
        summary = merge.apply_entries([winner], workbook_path=wb, trips_path=trips)
        assert summary["applied"] == 1
        assert summary["snapshots"], "snapshot must exist before apply"
        check = openpyxl.load_workbook(wb, read_only=True)
        assert check["Destinations"]["B2"].value == "resolved"
        check.close()
        merge.save_last_applied([winner], jdir)
        last = merge.load_last_applied(jdir)
        plan = merge.plan_merge([winner], [winner], last)
        assert plan["apply"] == [] and plan["conflicts"] == []


def test_offline_accumulation_queues_all_keys():
    local = [
        _entry("workbook", ["A", "Comment"], "1", "2026-09-27T10:00:00+00:00", "local"),
        _entry("workbook", ["B", "Comment"], "2", "2026-09-27T10:01:00+00:00", "local"),
        _entry("trips", ["trip-1", "variant-1"], {"n": 1}, "2026-09-27T10:02:00+00:00", "local"),
    ]
    plan = merge.plan_merge(local, [], {})
    assert len(plan["apply"]) == 3


def test_trips_delete_applies():
    with tempfile.TemporaryDirectory() as td:
        trips = Path(td) / "trips.json"
        data = _make_trips(trips)
        assert len(data["trips"][0]["variants"]) == 1
        delete = _entry(
            "trips", ["trip-1", "variant-1"], None,
            "2026-09-27T12:00:00+00:00", "cloud", op="delete",
        )
        summary = merge.apply_entries([delete], trips_path=trips)
        assert summary["applied"] == 1
        assert storage.load_from(trips)["trips"] == []


if __name__ == "__main__":
    test_newer_wins_without_conflict()
    print("PASS test_newer_wins_without_conflict")
    test_conflict_when_both_sides_changed()
    print("PASS test_conflict_when_both_sides_changed")
    test_resolution_and_idempotent_resync()
    print("PASS test_resolution_and_idempotent_resync")
    test_offline_accumulation_queues_all_keys()
    print("PASS test_offline_accumulation_queues_all_keys")
    test_trips_delete_applies()
    print("PASS test_trips_delete_applies")
    print("5/5 passed")
