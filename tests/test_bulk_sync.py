"""F13: the bulk data producers must journal their writes for sync.

Climate, AQI, cost, safety and food columns are written by scripts that go
straight through ``load_workbook_for_update`` / ``save_workbook_atomic``. Those
writes used to produce **no journal entry at all**, which is why the phone never
showed them. These tests pin the journal contract of each producer and the
"skipped" reporting for entries a device cannot apply.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import data_utils  # noqa: E402
from sync import journal, merge  # noqa: E402

CLIMATE_HEADERS = [
    "Destination", "Country", "Reviews", "Comment",
    "Jan High (C)", "Jan Low (C)", "Jan Rain (mm)", "Jan Rainy Days", "Jan AQI",
    "Feb High (C)", "Malaria risk?", "Avg AQI", "AQI Source", "Data Status",
]


@pytest.fixture()
def climate_book(tmp_path):
    path = tmp_path / "wb.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    ws.append(CLIMATE_HEADERS)
    ws.append(["Naples", "Italy", None, None] + [None] * (len(CLIMATE_HEADERS) - 4))
    ws.append(["Lima", "Peru", None, None] + [None] * (len(CLIMATE_HEADERS) - 4))
    other = wb.create_sheet("Airlines")
    other.append(["Airline", "Color"])
    wb.save(path)
    wb.close()
    return path


# ── the journal helper itself ────────────────────────────────────────────────

def test_journal_cell_changes_records_one_entry_per_cell(sandbox, monkeypatch):
    live = sandbox / "live.xlsx"
    live.write_bytes(b"x")
    monkeypatch.setattr(data_utils, "DATA_PATH", live)
    monkeypatch.setattr(journal, "DEFAULT_JOURNAL_DIR", sandbox / "journals")

    count = data_utils.journal_cell_changes(
        "Naples", {"Jan High (C)": 18.5, "Jan Low (C)": 9.0, "Comment": "hi"},
        live)
    assert count == 3
    entries = journal.read_entries(sandbox / "journals")
    assert {tuple(e["key"]) for e in entries} == {
        ("Naples", "Jan High (C)"), ("Naples", "Jan Low (C)"),
        ("Naples", "Comment")}
    assert all(e["store"] == "workbook" for e in entries)


def test_journal_cell_changes_ignores_empty_and_non_live(sandbox, monkeypatch):
    live = sandbox / "live.xlsx"
    live.write_bytes(b"x")
    other = sandbox / "elsewhere.xlsx"
    other.write_bytes(b"x")
    monkeypatch.setattr(data_utils, "DATA_PATH", live)
    # A path other than the live workbook is not journaled (test/script paths).
    assert data_utils.journal_cell_changes("Naples", {"A": 1}, other) == 0
    # ... and an empty change set is a no-op even for the live file.
    assert data_utils.journal_cell_changes("Naples", {}, live) == 0


def test_journal_cell_changes_is_silent_during_a_sync_apply(sandbox, monkeypatch):
    import os

    live = sandbox / "live.xlsx"
    live.write_bytes(b"x")
    monkeypatch.setattr(data_utils, "DATA_PATH", live)
    monkeypatch.setenv("SYNC_MERGE_APPLY", "1")
    assert data_utils.journal_cell_changes("Naples", {"A": 1}, live) == 0
        # (the env var is what stops a sync apply from re-journaling)


def test_journal_cell_changes_never_raises(sandbox, monkeypatch):
    live = sandbox / "live.xlsx"
    live.write_bytes(b"x")
    monkeypatch.setattr(data_utils, "DATA_PATH", live)

    def boom(*_a, **_k):
        raise RuntimeError("journal backend down")

    monkeypatch.setattr(journal, "ensure_baseline", boom)
    assert data_utils.journal_cell_changes("Naples", {"A": 1}, live) == 0


# ── aqi_api.update_destination_climate ──────────────────────────────────────

def test_climate_write_is_journalled(sandbox, monkeypatch, climate_book):
    import aqi_api

    monkeypatch.setattr(data_utils, "DATA_PATH", climate_book)
    monkeypatch.setattr(journal, "DEFAULT_JOURNAL_DIR", sandbox / "journals")

    message = aqi_api.update_destination_climate(
        "Naples",
        model_aqi={"Jan": 42, "Feb": 55},
        climate={"High (C)": [18.5] + [None] * 11,
                 "Low (C)": [9.0] + [None] * 11,
                 "Rain (mm)": [80.0] + [None] * 11,
                 "Rainy Days": [11.0] + [None] * 11},
        ground_aqi={"Jan": 40},
        model_typical=None, ground_typical=None,
        path=climate_book,
    )
    assert "Naples" in message

    entries = journal.read_entries(sandbox / "journals")
    written = {e["key"][1]: e["value"] for e in entries if e["key"][0] == "Naples"}
    assert written["Jan High (C)"] == 18.5
    assert written["Jan Low (C)"] == 9.0
    assert written["Jan AQI"] == 42
    assert written["Jan AQI (Ground)"] == 40
    assert "AQI Source" in written and "CAMS" in str(written["AQI Source"])
    # Nothing is journalled for months without data.
    assert not any(name.startswith("Feb High") for name in written)


def test_empty_climate_payload_leaves_the_workbook_alone(sandbox, monkeypatch,
                                                        climate_book):
    """A failed fetch used to overwrite the AQI Source provenance anyway."""
    import aqi_api

    before = climate_book.read_bytes()
    monkeypatch.setattr(data_utils, "DATA_PATH", climate_book)
    monkeypatch.setattr(journal, "DEFAULT_JOURNAL_DIR", sandbox / "journals")

    message = aqi_api.update_destination_climate(
        "Naples", model_aqi={}, climate={}, ground_aqi=None,
        path=climate_book)
    assert "nothing to write" in message
    assert climate_book.read_bytes() == before
    assert journal.read_entries(sandbox / "journals") == []


def test_aqi_source_string_has_balanced_parentheses(sandbox, monkeypatch,
                                                   climate_book):
    import aqi_api

    monkeypatch.setattr(data_utils, "DATA_PATH", climate_book)
    monkeypatch.setattr(journal, "DEFAULT_JOURNAL_DIR", sandbox / "journals")
    aqi_api.update_destination_climate(
        "Lima", model_aqi={"Jan": 30}, climate={}, ground_aqi=None,
        path=climate_book)
    wb = openpyxl.load_workbook(climate_book)
    ws = wb["Result sheet"]
    src = next(ws.cell(r, CLIMATE_HEADERS.index("AQI Source") + 1).value
               for r in (2, 3) if ws.cell(r, 1).value == "Lima")
    wb.close()
    assert src.count("(") == src.count(")"), src


def test_climate_write_reports_not_found_without_writing(sandbox, monkeypatch,
                                                         climate_book):
    import aqi_api

    before = climate_book.read_bytes()
    monkeypatch.setattr(data_utils, "DATA_PATH", climate_book)
    message = aqi_api.update_destination_climate(
        "Atlantis", model_aqi={"Jan": 1}, climate={}, ground_aqi=None,
        path=climate_book)
    assert "not found" in message
    assert climate_book.read_bytes() == before


# ── the apply side: skipped entries are reported, not dropped silently ──────

def _entry(destination, column, value, ts="2026-09-27T12:00:00+00:00",
           device="cloud"):
    return {"ts": ts, "device": device, "store": "workbook", "op": "upsert",
            "key": [destination, column], "value": value}


def test_apply_reports_a_missing_destination_row(sandbox, climate_book):
    summary = merge.apply_entries(
        [_entry("Naples", "Jan High (C)", 20.0),
         _entry("Atlantis", "Jan High (C)", 30.0)],
        workbook_path=climate_book, trips_path=sandbox / "trips.json")
    assert summary["applied"] == 1
    assert any("Atlantis" in line and "row not in this workbook" in line
               for line in summary["skipped"])


def test_apply_reports_a_missing_column(sandbox, climate_book):
    summary = merge.apply_entries(
        [_entry("Naples", "No Such Column", 1)],
        workbook_path=climate_book, trips_path=sandbox / "trips.json")
    assert summary["applied"] == 0
    assert any("no such column" in line.lower() for line in summary["skipped"])


def test_apply_lists_a_missing_destination_only_once(sandbox, climate_book):
    summary = merge.apply_entries(
        [_entry("Atlantis", "Jan High (C)", 1, ts="2026-09-27T12:00:00+00:00"),
         _entry("Atlantis", "Jan Low (C)", 2, ts="2026-09-27T12:01:00+00:00")],
        workbook_path=climate_book, trips_path=sandbox / "trips.json")
    assert len([line for line in summary["skipped"] if "Atlantis" in line]) == 1


def test_apply_reports_malformed_keys(sandbox, climate_book):
    bad = {"ts": "2026-09-27T12:00:00+00:00", "device": "cloud",
           "store": "workbook", "op": "upsert", "key": ["Naples"], "value": 1}
    summary = merge.apply_entries([bad], workbook_path=climate_book,
                                  trips_path=sandbox / "trips.json")
    assert any("malformed key" in line for line in summary["skipped"])


def test_apply_of_bulk_cells_writes_them_all(sandbox, climate_book):
    entries = [_entry("Naples", f"Jan {part}", i, ts=f"2026-09-27T12:0{i}:00+00:00")
               for i, part in enumerate(["High (C)", "Low (C)", "Rain (mm)"])]
    summary = merge.apply_entries(entries, workbook_path=climate_book,
                                  trips_path=sandbox / "trips.json")
    assert summary["applied"] == 3 and summary["skipped"] == []
    wb = openpyxl.load_workbook(climate_book)
    ws = wb["Result sheet"]
    assert ws.cell(2, CLIMATE_HEADERS.index("Jan High (C)") + 1).value == 0
    assert ws.cell(2, CLIMATE_HEADERS.index("Jan Rain (mm)") + 1).value == 2
    wb.close()


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit("This module needs pytest: python -m pytest tests/test_bulk_sync.py")
