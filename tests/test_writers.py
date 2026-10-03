"""Contract tests for the workbook writers.

Two regressions are covered:

* **Silent no-ops.** Every writer used to ``return`` without writing when it
  could not find the row or the column, and the UI treated "no exception" as
  success — so clicking the heart / ❔ / Visited / Prio button could do nothing
  at all while the page claimed it worked.
* **Writer/reader disagreement.** The reader accepted five aliases for the
  favourite column but the writer only understood ``auswahl``, so a workbook
  the UI could read was not the workbook the UI wrote.

Needs pytest: ``python -m pytest tests/test_writers.py``
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import data_utils  # noqa: E402

HEADERS = ["Destination", "Country", "Continent", "Comment", "Visited?",
           "In näherer Auswahl 2025?", "Prio Thorsten", "To be researched",
           "Food - Spicyness", "Food - Description", "Food - Main Dishes",
           "Data Status"]


@pytest.fixture()
def book(tmp_path):
    """A minimal workbook with the real header names, plus a second sheet."""
    path = tmp_path / "wb.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    ws.append(HEADERS)
    ws.append(["Naples", "Italy", None, None, None, None, None, None, None, None])
    ws.append(["Oslo", "Norway", None, None, None, None, None, None, None, None])
    other = wb.create_sheet("Airlines")
    other.append(["Airline", "Color"])
    wb.save(path)
    wb.close()
    return path


def _cell(path: Path, header: str, row: int = 2):
    """Read one cell by *header name* (works for any workbook shape)."""
    wb = openpyxl.load_workbook(path)
    ws = wb["Result sheet"]
    idx = next(i for i in range(1, ws.max_column + 1)
               if str(ws.cell(row=1, column=i).value).strip() == header)
    value = ws.cell(row=row, column=idx).value
    wb.close()
    return value


# ── happy path: every writer reports success and writes its cell ─────────────

def test_update_favorite_writes_and_reports_success(book):
    assert data_utils.update_favorite_status("Naples", True, path=book) is True
    assert _cell(book, "In näherer Auswahl 2025?") == "x"
    assert data_utils.update_favorite_status("Naples", False, path=book) is True
    assert _cell(book, "In näherer Auswahl 2025?") is None


def test_update_visited_writes_and_reports_success(book):
    assert data_utils.update_visited_status("Naples", True, path=book) is True
    assert _cell(book, "Visited?") is True


def test_update_research_writes_and_reports_success(book):
    assert data_utils.update_to_be_researched_status("Naples", True, path=book) is True
    assert _cell(book, "To be researched") is True


def test_update_prio_writes_number_or_clears(book):
    assert data_utils.update_prio_thorsten("Naples", 7, path=book) is True
    assert _cell(book, "Prio Thorsten") == 7
    assert data_utils.update_prio_thorsten("Naples", None, path=book) is True
    assert _cell(book, "Prio Thorsten") is None


def test_update_prio_accepts_a_numeric_string(book):
    assert data_utils.update_prio_thorsten("Naples", "8", path=book) is True
    assert _cell(book, "Prio Thorsten") == 8


def test_update_prio_unparseable_text_clears_the_cell(book):
    assert data_utils.update_prio_thorsten("Naples", "not a number", path=book) is True
    assert _cell(book, "Prio Thorsten") is None


def test_update_comment_writes_and_clears(book):
    assert data_utils.update_comment("Naples", "great pizza", path=book) is True
    assert _cell(book, "Comment") == "great pizza"
    assert data_utils.update_comment("Naples", "   ", path=book) is True
    assert _cell(book, "Comment") is None


def test_update_food_writes_all_three_columns(book):
    assert data_utils.update_food("Naples", 6, "tasty", ["pizza", "sorbet"],
                                  path=book) is True
    assert _cell(book, "Food - Spicyness") == 6.0
    assert _cell(book, "Food - Description") == "tasty"
    assert _cell(book, "Food - Main Dishes") == "pizza, sorbet"


def test_update_reviews_needs_its_columns(tmp_path):
    path = tmp_path / "slim.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    ws.append(["Destination", "Reviews"])
    ws.append(["Naples", 5.0])
    wb.save(path)
    wb.close()
    # Only one of the four review columns exists -> still writes that one.
    assert data_utils.update_reviews("Naples", 8.5, "text", "p", "d", path=path) is True
    assert _read_slim(path) == 8.5


def _read_slim(path: Path):
    wb = openpyxl.load_workbook(path)
    value = wb["Result sheet"].cell(row=2, column=2).value
    wb.close()
    return value


# ── missing row / column must be reported, never silently ignored ────────────

WRITERS = [
    ("favorite", lambda p: data_utils.update_favorite_status("Atlantis", True, path=p)),
    ("visited", lambda p: data_utils.update_visited_status("Atlantis", True, path=p)),
    ("research", lambda p: data_utils.update_to_be_researched_status("Atlantis", True, path=p)),
    ("prio", lambda p: data_utils.update_prio_thorsten("Atlantis", 3, path=p)),
    ("comment", lambda p: data_utils.update_comment("Atlantis", "x", path=p)),
    ("food", lambda p: data_utils.update_food("Atlantis", 5, "y", ["z"], path=p)),
    ("reviews", lambda p: data_utils.update_reviews("Atlantis", 7.0, "a", "b", "c", path=p)),
]


@pytest.mark.parametrize("name,writer", WRITERS, ids=[w[0] for w in WRITERS])
def test_unknown_destination_reports_nothing_written(book, name, writer):
    assert writer(book) is False, f"{name} must report False for a missing row"


def test_missing_target_column_reports_nothing_written(tmp_path):
    path = tmp_path / "sparse.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    ws.append(["Destination", "Country"])       # no visited/prio/comment columns
    ws.append(["Naples", "Italy"])
    wb.save(path)
    wb.close()
    assert data_utils.update_visited_status("Naples", True, path=path) is False
    assert data_utils.update_prio_thorsten("Naples", 5, path=path) is False
    assert data_utils.update_favorite_status("Naples", True, path=path) is False


def test_missing_destination_column_reports_nothing_written(tmp_path):
    # NOTE: "Continent" is included in both mini-workbooks only so that
    # _find_destination_sheet() recognises the sheet — sheet detection keys off
    # those column names, independently of the alias sets.
    # 'City' is a recognised destination alias -> writable.
    aliased = tmp_path / "aliased.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    ws.append(["City", "Continent", "Visited?"])
    ws.append(["Naples", "Europe", None])
    wb.save(aliased)
    wb.close()
    assert data_utils.update_visited_status("Naples", True, path=aliased) is True

    # ... but a sheet without any recognisable destination column must fail loudly.
    path = tmp_path / "nodest.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result sheet"
    ws.append(["Foo", "Continent", "Visited?"])
    ws.append(["Naples", "Europe", None])
    wb.save(path)
    wb.close()
    assert data_utils.update_visited_status("Naples", True, path=path) is False


def test_no_destination_sheet_reports_nothing_written(tmp_path):
    path = tmp_path / "nosheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SomeOtherSheet"
    ws.append(["Foo", "Bar"])
    ws.append(["Naples", 1])
    wb.save(path)
    wb.close()
    assert data_utils.update_comment("Naples", "x", path=path) is False


def test_writer_never_touches_the_other_sheets(book):
    before = _airlines_rows(book)
    data_utils.update_comment("Naples", "hello", path=book)
    assert _airlines_rows(book) == before


def _airlines_rows(path: Path) -> list[list]:
    wb = openpyxl.load_workbook(path)
    ws = wb["Airlines"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    return rows


# ── reader/writer use the same alias vocabulary ─────────────────────────────

def test_match_header_index_prefers_an_exact_match():
    headers = {"In näherer Auswahl 2025?": 2, "Auswahl Kommentar": 3}
    assert data_utils.match_header_index(headers, data_utils.NEARER_ALIASES) == 2


def test_match_header_index_falls_back_to_a_substring():
    headers = {"Nähere Auswahl alt": 4}
    assert data_utils.match_header_index(headers, data_utils.NEARER_ALIASES) == 4


def test_match_header_index_returns_none_when_absent():
    assert data_utils.match_header_index({"Foo": 1}, data_utils.VISITED_ALIASES) is None


def test_english_favourite_header_is_writable():
    """A workbook that says 'Nearer' must be writable, not read-only."""
    import tempfile

    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "wb.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Result sheet"
        ws.append(["Destination", "Nearer", "Visited?"])
        ws.append(["Naples", None, None])
        wb.save(path)
        wb.close()
        assert data_utils.update_favorite_status("Naples", True, path=path) is True
        assert _cell(path, "Nearer") == "x"


# ── add_new_destination keeps its (success, message) contract ───────────────

def test_add_new_destination_rejects_duplicates(book):
    ok, message = data_utils.add_new_destination("Naples", "Italy", "Europe", path=book)
    assert ok is False and "already exists" in message


def test_add_new_destination_rejects_an_empty_name(book):
    ok, message = data_utils.add_new_destination("   ", "Italy", "Europe", path=book)
    assert ok is False and "empty" in message.lower()


def test_add_new_destination_appends_with_placeholders(book):
    ok, _message = data_utils.add_new_destination("Porto", "Portugal", "Europe", path=book)
    assert ok is True
    assert _cell(book, "Destination", row=4) == "Porto"
    assert _cell(book, "Country", row=4) == "Portugal"
    assert _cell(book, "Continent", row=4) == "Europe"
    assert "PLACEHOLDER" in str(_cell(book, "Data Status", row=4))


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit("This module needs pytest: python -m pytest tests/test_writers.py")
